"""Тесты генерации и выгрузки печатного бланка тестирования Word (.docx)."""

import io
from datetime import timedelta
import docx
from django.contrib.auth.models import Group
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone

from customers_app.models import DataBaseUser, Job, Division, DataBaseUserWorkProfile
from testing_app.models import (
    Testing,
    QuestionCategory,
    Question,
    AnswerOption,
    TestingGroup,
    TestingCategorySetting,
    TestingAssignment,
    TestingAttempt,
    AttemptQuestion,
)
from testing_app.services.blank_service import (
    get_questions_for_blank,
    append_questions_sheet_to_document,
    generate_filled_testing_blank_bytes,
    generate_event_testing_blank_bytes,
    send_testing_blank_by_email,
)


class BlankServiceTests(TestCase):
    """Набор тестов для проверки генератора документов Word (.docx) бланков тестирования."""

    def setUp(self) -> None:
        """Подготовка тестовых пользователей, мероприятия, категорий и вопросов."""
        self.manager_group, _ = Group.objects.get_or_create(name="Ответственные за тестирование")

        self.manager_user = DataBaseUser.objects.create_user(
            username="manager_test",
            email="manager@barkol.ru",
            password="testpassword123",
            first_name="Пётр",
            last_name="Менеджеров",
        )
        self.manager_user.groups.add(self.manager_group)

        self.employee_user = DataBaseUser.objects.create_user(
            username="employee_test",
            email="employee@barkol.ru",
            password="testpassword123",
            first_name="Иван",
            last_name="Иванов",
            surname="Иванович",
        )
        self.job = Job.objects.create(name="Авиатехник по ПиД")
        self.division = Division.objects.create(name="АТБ Тверь")
        self.employee_user.user_work_profile = DataBaseUserWorkProfile.objects.create(
            job=self.job,
            divisions=self.division,
        )
        self.employee_user.save(update_fields=["user_work_profile"])

        now = timezone.now()
        self.testing = Testing.objects.create(
            title="Периодическая подготовка ИТП по ТО ВС",
            order_number="45-ТО",
            order_date=now.date(),
            order_name="О проведении аттестации",
            start_datetime=now - timedelta(days=1),
            end_datetime=now + timedelta(days=10),
            questions_count=5,
            passing_score_percentage=80,
            status=Testing.Status.ACTIVE,
        )

        self.group_performing = TestingGroup.objects.create(
            testing=self.testing,
            name="Выполняющие ТО ВС",
            code=TestingGroup.Code.PERFORMING,
        )
        self.group_ensuring = TestingGroup.objects.create(
            testing=self.testing,
            name="Обеспечение ТО ВС",
            code=TestingGroup.Code.ENSURING,
        )

        self.category1 = QuestionCategory.objects.create(name="Воздушное законодательство", is_active=True)
        self.category2 = QuestionCategory.objects.create(name="СУБП и безопасность полетов", is_active=True)

        TestingCategorySetting.objects.create(
            testing=self.testing,
            category=self.category1,
            percentage=60,
            calculated_questions_count=3,
        )
        TestingCategorySetting.objects.create(
            testing=self.testing,
            category=self.category2,
            percentage=40,
            calculated_questions_count=2,
        )

        # Создаем банк вопросов: по 3 вопроса в каждой категории
        self.questions = []
        for i in range(3):
            q1 = Question.objects.create(
                category=self.category1,
                text=f"Вопрос по законодательству №{i + 1}",
                explanation=f"Пояснение к вопросу №{i + 1}",
                status=Question.Status.ACTIVE,
            )
            AnswerOption.objects.create(question=q1, text="Правильный ответ", order_num=1, is_correct=True)
            AnswerOption.objects.create(question=q1, text="Неверный ответ", order_num=2, is_correct=False)
            self.questions.append(q1)

            q2 = Question.objects.create(
                category=self.category2,
                text=f"Вопрос по СУБП №{i + 1}",
                explanation=f"Пояснение к СУБП №{i + 1}",
                status=Question.Status.ACTIVE,
            )
            AnswerOption.objects.create(question=q2, text="Ответ СУБП верный", order_num=1, is_correct=True)
            AnswerOption.objects.create(question=q2, text="Ответ СУБП неверный", order_num=2, is_correct=False)
            self.questions.append(q2)

        self.assignment = TestingAssignment.objects.create(
            testing=self.testing,
            group=self.group_performing,
            employee=self.employee_user,
            assigned_job_title=self.job.name,
            assigned_division_title=self.division.name,
            status=TestingAssignment.Status.NOT_STARTED,
        )

        self.client = Client()

    def test_get_questions_for_blank_from_category_settings(self) -> None:
        """Проверяет отбор вопросов по квотам категорий, если попытка еще не создана."""
        questions_data = get_questions_for_blank(self.testing, assignment=self.assignment)
        self.assertEqual(len(questions_data), 5)
        # 3 вопроса из category1 и 2 из category2
        cat1_count = sum(1 for q in questions_data if q["category_name"] == self.category1.name)
        cat2_count = sum(1 for q in questions_data if q["category_name"] == self.category2.name)
        self.assertEqual(cat1_count, 3)
        self.assertEqual(cat2_count, 2)
        # Проверяем наличие вариантов ответов
        for q in questions_data:
            self.assertEqual(len(q["options"]), 2)
            self.assertTrue(any(opt["is_correct"] for opt in q["options"]))

    def test_get_questions_for_blank_from_attempt_snapshot(self) -> None:
        """Проверяет, что при наличии попытки вопросы берутся из неизменяемого снимка AttemptQuestion."""
        attempt = TestingAttempt.objects.create(
            assignment=self.assignment,
            attempt_number=1,
            planned_end_at=timezone.now() + timedelta(hours=1),
            total_questions=2,
        )
        aq = AttemptQuestion.objects.create(
            attempt=attempt,
            source_question=self.questions[0],
            category_name="Кастомный снимок категории",
            order_num=1,
            question_text="Зафиксированный текст вопроса из снимка",
            options_snapshot=[
                {"id": 1, "order_num": 1, "text": "Опция А", "is_correct": False},
                {"id": 2, "order_num": 2, "text": "Опция Б", "is_correct": True},
            ],
        )

        questions_data = get_questions_for_blank(self.testing, assignment=self.assignment)
        self.assertEqual(len(questions_data), 1)
        self.assertEqual(questions_data[0]["order_num"], 1)
        self.assertEqual(questions_data[0]["text"], "Зафиксированный текст вопроса из снимка")
        self.assertEqual(questions_data[0]["category_name"], "Кастомный снимок категории")
        self.assertEqual(questions_data[0]["options"][1]["text"], "Опция Б")
        self.assertTrue(questions_data[0]["options"][1]["is_correct"])

    def test_generate_filled_testing_blank_bytes(self) -> None:
        """Проверяет формирование DOCX-бланка для назначения сотрудника."""
        docx_bytes, filename = generate_filled_testing_blank_bytes(
            self.assignment,
            user=self.manager_user,
            with_answers=False,
        )
        self.assertIsInstance(docx_bytes, bytes)
        self.assertGreater(len(docx_bytes), 5000)
        self.assertTrue(filename.endswith(".docx"))
        self.assertIn("Иванов", filename)

        # Открываем сгенерированный Word документ и проверяем наличие ключевых строк
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                full_text += "\n" + " ".join(c.text for c in row.cells)

        self.assertIn("ЛИСТ С ВОПРОСАМИ ИТОГОВОГО ТЕСТИРОВАНИЯ", full_text)
        self.assertIn("Иванов Иван Иванович", full_text)
        self.assertIn("Приказ № 45-ТО", full_text)
        self.assertIn("РЕЗУЛЬТАТ ПРОВЕРКИ И ОЦЕНКА КОМИССИИ", full_text)
        self.assertIn("☐  ЗАЧТЕНО", full_text)

    def test_generate_testing_blank_with_answers_key(self) -> None:
        """Проверяет формирование ключа с ответами для экзаменационной комиссии."""
        docx_bytes, filename = generate_filled_testing_blank_bytes(
            self.assignment,
            user=self.manager_user,
            with_answers=True,
        )
        self.assertIn("Ключ_ответов", filename)

        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("КЛЮЧ С ПРАВИЛЬНЫМИ ОТВЕТАМИ", full_text)
        self.assertIn("ПРАВИЛЬНЫЙ ОТВЕТ", full_text)
        self.assertIn("☑", full_text)

    def test_generate_event_testing_blank_bytes(self) -> None:
        """Проверяет формирование общего бланка мероприятия без привязки к сотруднику."""
        docx_bytes, filename = generate_event_testing_blank_bytes(
            testing=self.testing,
            group=self.group_performing,
            user=self.manager_user,
            with_answers=False,
        )
        self.assertIsInstance(docx_bytes, bytes)
        self.assertGreater(len(docx_bytes), 5000)
        self.assertIn("Бланк_тестирования_Приказ_45-ТО", filename)

        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                full_text += "\n" + " ".join(c.text for c in row.cells)

        self.assertIn("ЛИСТ С ВОПРОСАМИ ИТОГОВОГО ТЕСТИРОВАНИЯ", full_text)
        self.assertIn("Приказ № 45-ТО", full_text)

    def test_download_blank_view_by_employee(self) -> None:
        """Проверяет скачивание бланка обычным сотрудником через HTTP View."""
        self.client.force_login(self.employee_user)
        url = reverse("testing_app:download_blank", kwargs={"assignment_id": self.assignment.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("attachment;", response["Content-Disposition"])

    def test_download_blank_view_blocks_answers_for_employee(self) -> None:
        """Проверяет, что обычному сотруднику запрещено скачивать бланк с ключом ответов."""
        self.client.force_login(self.employee_user)
        url = reverse("testing_app:download_blank", kwargs={"assignment_id": self.assignment.id}) + "?with_answers=1"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        # Должен вернуться обычный бланк без правильных ответов
        doc = docx.Document(io.BytesIO(response.content))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("ПРАВИЛЬНЫЙ ОТВЕТ", full_text)

    def test_download_blank_view_allows_answers_for_manager(self) -> None:
        """Проверяет, что менеджеру тестирования разрешено скачивать ключ ответов."""
        self.client.force_login(self.manager_user)
        url = reverse("testing_app:download_blank", kwargs={"assignment_id": self.assignment.id}) + "?with_answers=1"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        doc = docx.Document(io.BytesIO(response.content))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("ПРАВИЛЬНЫЙ ОТВЕТ", full_text)

    def test_download_event_blank_view_by_manager(self) -> None:
        """Проверяет скачивание общего бланка мероприятия менеджером."""
        self.client.force_login(self.manager_user)
        url = reverse("testing_app:event_blank_download", kwargs={"pk": self.testing.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def test_download_event_blank_view_forbidden_for_employee(self) -> None:
        """Проверяет, что обычному сотруднику запрещен доступ к выгрузке бланков мероприятия."""
        self.client.force_login(self.employee_user)
        url = reverse("testing_app:event_blank_download", kwargs={"pk": self.testing.id})
        response = self.client.get(url)
        # TestingManagerRequiredMixin редиректит на testing_app:my_tests
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("testing_app:my_tests"), response["Location"])
